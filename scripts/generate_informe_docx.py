#!/usr/bin/env python3
"""Genera el Informe Técnico (.docx) del simulador de redes."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Informe_Tecnico_Simulador_Redes.docx"
SHOTS = ROOT / "docs" / "screenshots"


def set_run_font(run, size=11, bold=False, italic=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True)
    return h


def add_para(doc, text, *, bold=False, italic=False, size=11, space_after=8):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=11)


def add_formula(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=12, italic=True)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                set_run_font(run, bold=True, size=10)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()


def add_image(doc, path: Path, caption: str, width_cm=14.5):
    if not path.exists():
        add_para(doc, f"[Captura no encontrada: {path.name}]", italic=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    set_run_font(r, size=10, italic=True, color=RGBColor(80, 80, 80))


def build() -> Path:
    doc = Document()

    # Márgenes
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # Portada
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("INFORME TÉCNICO")
    set_run_font(r, size=22, bold=True)

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run("Simulador Dinámico de Redes de Computadoras")
    set_run_font(r, size=16, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(
        "Integración de Teoría de Colas, Modelos de Inventario\n"
        "y Algoritmo Húngaro de Asignación Óptima"
    )
    set_run_font(r, size=12, italic=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        "Métodos Cuantitativos\n"
        "Lenguaje: Python  ·  Pygame + SimPy + NumPy/SciPy + Gemini API\n"
        "Branch: feature/simulador-red-dinamico"
    )
    set_run_font(r, size=11)

    doc.add_page_break()

    # 1. Introducción
    add_heading(doc, "1. Introducción y objetivo", 1)
    add_para(
        doc,
        "El presente informe documenta el diseño, modelado matemático e implementación de un "
        "simulador visual e interactivo de una red de computadoras. El sistema integra tres "
        "pilares de los métodos cuantitativos aplicados al dominio de redes:",
    )
    add_bullets(
        doc,
        [
            "Teoría de líneas de espera (colas) para el tráfico de paquetes y buffers.",
            "Modelos de inventario para la capacidad de cola y el control de flujo.",
            "Modelo de asignación (Algoritmo Húngaro) para el enrutamiento dinámico óptimo.",
        ],
    )
    add_para(
        doc,
        "Adicionalmente, el software registra métricas de desempeño, las exporta a un archivo "
        "plano (reporte_simulacion.txt) y las envía a la API de Gemini para obtener un "
        "diagnóstico automatizado con conclusiones y recomendaciones de optimización.",
    )

    # 2. Arquitectura
    add_heading(doc, "2. Arquitectura del software", 1)
    add_para(
        doc,
        "La solución se organizó en módulos desacoplados bajo src/simulator/, con main.py como "
        "punto de entrada (modo GUI o headless):",
    )
    add_table(
        doc,
        ["Módulo", "Archivo", "Responsabilidad"],
        [
            ["A — Colas", "queueing.py", "Poisson(λ), Exp(μ), métricas L, Lq, W, Wq"],
            ["B — Inventario", "inventory.py", "Buffer S, política (s, Q), costos holding/shortage"],
            ["C — Asignación", "assignment.py", "Matriz Cᵢⱼ y Algoritmo Húngaro (SciPy)"],
            ["Red", "network.py", "Nodos, enlaces, paquetes y topología"],
            ["Motor", "simulation.py", "Procesos discretos SimPy acoplados al reloj"],
            ["GUI", "gui.py", "Visualización e interacción Pygame"],
            ["Reporte", "report.py", "Exportación estructurada a TXT"],
            ["API", "api_client.py", "Análisis con Gemini (Interactions API)"],
        ],
    )
    add_para(
        doc,
        "La topología simulada consta de cinco nodos: Fuente → {R1, R2, R3} → Destino, con "
        "enlaces entre routers vecinos. Cada nodo transmisor actúa como servidor con tasa μ "
        "y un buffer de inventario de capacidad S.",
    )

    # 3. Modelos matemáticos
    add_heading(doc, "3. Modelos matemáticos", 1)

    add_heading(doc, "3.1 Teoría de colas (líneas de espera)", 2)
    add_para(
        doc,
        "Las llegadas de paquetes a la Fuente siguen un proceso de Poisson con tasa λ "
        "configurable. El tiempo entre llegadas se genera como variable exponencial de media 1/λ. "
        "Cada nodo transmisor atiende con tiempo de servicio exponencial de media 1/μ "
        "(equivalente a un servidor M/M/1 por nodo, con capacidad finita dada por el buffer).",
    )
    add_para(doc, "Métricas calculadas en tiempo real (estimadores empíricos / Little):", bold=True)
    add_bullets(
        doc,
        [
            "L: número promedio de paquetes en el sistema (∫ N(t) dt / T).",
            "Lq: número promedio de paquetes en cola (∫ Nq(t) dt / T).",
            "W: tiempo medio de estancia extremo a extremo (promedio de sojourn times).",
            "Wq: tiempo medio de espera en cola antes del primer servicio.",
        ],
    )
    add_formula(doc, "ρ = λ / μ    (factor de utilización; se requiere ρ < 1 para estabilidad)")
    add_para(
        doc,
        "Con los parámetros por defecto λ = 15 paquetes/s y μ = 18 paquetes/s se obtiene "
        "ρ ≈ 0.833, es decir, un régimen estable con ocupación moderada-alta.",
    )

    add_heading(doc, "3.2 Modelo de inventario (control de buffer)", 2)
    add_para(
        doc,
        "Cada router gestiona un inventario de paquetes en memoria (buffer) con los siguientes "
        "parámetros de política de revisión continua inspirada en (s, Q):",
    )
    add_bullets(
        doc,
        [
            "S: capacidad máxima del buffer (inventario).",
            "s: umbral de reabastecimiento / control de flujo.",
            "Q: tamaño del lote (créditos de admisión) solicitado al cruzar por debajo de s.",
        ],
    )
    add_para(
        doc,
        "Cuando el nivel del buffer cae por debajo de s, el nodo emite una señal de control de "
        "flujo y otorga hasta Q créditos de admisión (reapertura del canal de entrada). Se aplica "
        "histéresis: la señal no se repite hasta que el nivel recupere un valor ≥ s.",
    )
    add_para(doc, "Costos del sistema:", bold=True)
    add_bullets(
        doc,
        [
            "Costo de mantenimiento (holding): proporcional a nivel × tiempo × tasa de holding "
            "(latencia/ocupación de RAM).",
            "Costo de ruptura (shortage): penalización fija por cada paquete descartado ante "
            "overflow (nivel = S).",
            "Costo global = holding + shortage.",
        ],
    )

    add_heading(doc, "3.3 Modelo de asignación (Algoritmo Húngaro)", 2)
    add_para(
        doc,
        "Cada intervalo discreto Δt el simulador construye una matriz de costos entre flujos "
        "pendientes y enlaces/nodos de salida disponibles:",
    )
    add_formula(doc, "Cᵢⱼ = Latencia_actual(i,j) + α · Saturación_buffer(j)")
    add_para(
        doc,
        "Donde la saturación es nivel/S ∈ [0,1] y α pondera la congestión frente a la latencia. "
        "La asignación uno-a-uno de mínimo costo se resuelve con el Algoritmo Húngaro "
        "(implementación SciPy: scipy.optimize.linear_sum_assignment). El resultado actualiza la "
        "tabla de ruteo (p. ej. Fuente → Rk) para balancear la carga dinámicamente. Enlaces "
        "fallidos se marcan con costo muy alto para excluirlos de la solución óptima.",
    )

    # 4. Interfaz
    add_heading(doc, "4. Interfaz gráfica (Pygame)", 1)
    add_para(doc, "Representación visual:", bold=True)
    add_bullets(
        doc,
        [
            "Nodos como círculos coloreados por saturación: verde (<50%), amarillo (50–80%), rojo (>80%).",
            "Enlaces como líneas dinámicas (rojo/tachado si fallan).",
            "Paquetes animados con progreso sincronizado al tiempo de tránsito SimPy; colas apiladas junto a cada nodo.",
            "Dashboard lateral: tiempo, λ, μ, S, s, Q, L/Lq/W/Wq, pérdidas, costos y corridas del Húngaro.",
        ],
    )
    add_para(doc, "Controles interactivos:", bold=True)
    add_bullets(
        doc,
        [
            "SPACE: pausar / reanudar.",
            "↑ / ↓: ajustar λ en tiempo real.",
            "F: simular falla de un enlace; R: reparar enlaces.",
            "[ / ]: cambiar velocidad de simulación.",
            "E: exportar reporte y consultar Gemini; ESC: exportar, analizar y salir.",
        ],
    )

    add_heading(doc, "4.1 Capturas de la interfaz en ejecución", 2)
    add_image(
        doc,
        SHOTS / "gui_inicio.png",
        "Figura 1. GUI en los primeros segundos: topología Fuente–R1/R2/R3–Destino y dashboard de métricas.",
    )
    add_image(
        doc,
        SHOTS / "gui_ejecucion.png",
        "Figura 2. GUI en régimen estable: paquetes en tránsito, colas por nodo y costos acumulados.",
    )

    # 5. Resultados
    add_heading(doc, "5. Resultados de la simulación", 1)
    add_para(
        doc,
        "A continuación se resumen los parámetros y métricas de una corrida representativa "
        "(exportada a reporte_simulacion.txt):",
    )
    add_table(
        doc,
        ["Parámetro / métrica", "Valor"],
        [
            ["Tiempo de simulación", "16.64 s"],
            ["λ (llegada)", "15.0 paquetes/s"],
            ["μ (servicio)", "18.0 paquetes/s"],
            ["S (capacidad buffer)", "50 paquetes"],
            ["s (umbral)", "10 paquetes"],
            ["Paquetes procesados", "213"],
            ["Paquetes perdidos (overflow)", "0"],
            ["Tasa de pérdida", "0.00%"],
            ["Wq (espera en cola)", "0.13 s"],
            ["W (estancia e2e)", "1.01 s"],
            ["L (promedio en sistema)", "13.26"],
            ["Lq (promedio en cola)", "3.54"],
            ["Costo holding", "$2.39"],
            ["Costo shortage", "$0.00"],
            ["Costo global", "$2.39"],
            ["Ejecuciones Algoritmo Húngaro", "16"],
            ["Señales de control de flujo", "3"],
        ],
    )
    add_para(
        doc,
        "Interpretación breve: con ρ ≈ 0.83 el sistema se mantiene estable y sin pérdidas. "
        "El costo global proviene íntegramente del holding (no hubo ruptura de inventario). "
        "El Algoritmo Húngaro se ejecutó aproximadamente una vez por segundo (Δt = 1 s) y "
        "ajustó el ruteo hacia el router menos saturado en cada intervalo.",
    )

    # 6. API
    add_heading(doc, "6. Integración con Gemini API", 1)
    add_para(
        doc,
        "Al finalizar la simulación (tecla E o ESC), el programa guarda el reporte TXT y envía "
        "el siguiente prompt a Gemini (modelo gemini-3.6-flash, Interactions API):",
    )
    add_para(
        doc,
        '"Analiza los siguientes resultados de desempeño de un simulador de red basado en '
        "teoría de colas e inventario. Evalúa la tasa de pérdida de paquetes, tiempos de espera "
        'y costos, e indica conclusiones detalladas y 3 recomendaciones de optimización."',
        italic=True,
    )
    add_para(
        doc,
        "La respuesta se imprime en consola y se anexa al final de reporte_simulacion.txt. "
        "En la corrida documentada, Gemini confirmó la estabilidad (cero pérdidas), señaló el "
        "overhead de latencia e2e respecto al modelo M/M/1 puro, y propuso tres recomendaciones: "
        "(1) reducir la frecuencia del Húngaro a un esquema reactivo, (2) reforzar la penalización "
        "por saturación de ruta en Cᵢⱼ, y (3) ajustar S y s para bajar el costo de almacenamiento "
        "sin comprometer la tasa de pérdida.",
    )

    # 7. Conclusiones
    add_heading(doc, "7. Conclusiones", 1)
    add_bullets(
        doc,
        [
            "Se cumplió el objetivo de integrar colas, inventario y asignación óptima en un "
            "simulador visual interactivo con exportación de métricas y análisis por API.",
            "El mapeo matemático → red es coherente: Poisson/Exp para tráfico, (s,Q) para buffers "
            "y Húngaro para balanceo dinámico de enlaces.",
            "En las condiciones base (λ=15, μ=18, S=50) el sistema opera estable, sin overflow y "
            "con costos bajos dominados por holding.",
            "La interfaz permite experimentar en vivo (cambiar λ, fallar enlaces, pausar) y "
            "observar el impacto inmediato sobre L, Wq, pérdidas y la tabla de ruteo.",
            "Como trabajo futuro: informe de sensibilidad (barrido de λ, S, α), políticas (s,S) "
            "alternativas y comparación contra enrutamiento estático sin Húngaro.",
        ],
    )

    add_heading(doc, "8. Referencias de implementación", 1)
    add_bullets(
        doc,
        [
            "SimPy — simulación de eventos discretos.",
            "Pygame — interfaz gráfica e interacción.",
            "NumPy / SciPy (linear_sum_assignment) — matriz de costos y Algoritmo Húngaro.",
            "google-genai — cliente de Gemini API.",
            "Repositorio: branch feature/simulador-red-dinamico (PR #1).",
        ],
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Generado: {path}")
